import os
import re
import json
import base64
from typing import List, Dict, Any, Optional, Tuple
import openpyxl
from cryptography.fernet import Fernet
from dotenv import load_dotenv


# ─── Startup Stage Constants ────────────────────────────────────────
STAGE_IDEATION = "Ideation"
STAGE_PROTOTYPE = "Prototype"
STAGE_MVP = "MVP/Pre-Revenue"
STAGE_REVENUE = "Revenue"
STAGE_GROWTH = "Growth/Scaling"
STAGE_SEED = "Seed Stage"

ALL_STAGES = [STAGE_IDEATION, STAGE_PROTOTYPE, STAGE_MVP, STAGE_REVENUE, STAGE_GROWTH, STAGE_SEED]

STAGE_COLORS = {
    STAGE_IDEATION:    {"bg": "#EEF2FF", "text": "#4338CA", "border": "#C7D2FE"},
    STAGE_PROTOTYPE:   {"bg": "#F0FDF4", "text": "#166534", "border": "#BBF7D0"},
    STAGE_MVP:         {"bg": "#FFF7ED", "text": "#9A3412", "border": "#FED7AA"},
    STAGE_REVENUE:     {"bg": "#FEF2F2", "text": "#991B1B", "border": "#FECACA"},
    STAGE_GROWTH:      {"bg": "#FDF4FF", "text": "#86198F", "border": "#F0ABFC"},
    STAGE_SEED:        {"bg": "#ECFDF5", "text": "#065F46", "border": "#A7F3D0"},
}

# ─── Stage Auto-Classification ──────────────────────────────────────

def _normalize_stage_hint(raw: str) -> Optional[str]:
    if not raw:
        return None
    s = raw.strip().lower()
    if not s or s in ["none", "na", "n/a", "nil", "active", "other"]:
        return None
    mapping = {
        STAGE_IDEATION:    ["idea", "ideation", "concept", "conceptual", "brainstorm", "pre-idea"],
        STAGE_PROTOTYPE:   ["prototype", "poc", "proof of concept", "proof-of-concept", "alpha"],
        STAGE_MVP:         ["mvp", "pre-revenue", "pre revenue", "beta", "early product", "product built"],
        STAGE_REVENUE:     ["revenue", "revenue generating", "early revenue", "monetizing", "earning"],
        STAGE_GROWTH:      ["growth", "scaling", "scale-up", "scale up", "expansion", "expanding"],
        STAGE_SEED:        ["seed", "seed stage", "seed round", "pre-series a", "fundraise", "fundraising"],
    }
    for canonical, keywords in mapping.items():
        for kw in keywords:
            if kw in s:
                return canonical
    return None


def classify_startup_stage(data: Dict[str, Any]) -> str:
    raw_stage = str(data.get("stage") or "").strip()
    hint = _normalize_stage_hint(raw_stage)
    if hint:
        return hint

    revenue = float(data.get("revenue") or 0.0)
    team_size = int(data.get("team_size") or 1)
    has_website = bool(str(data.get("website") or "").strip() and str(data.get("website")).strip().startswith("http"))
    has_dpiit = bool(data.get("dpiit", False))
    has_pitch = bool(str(data.get("pitch_deck_url") or "").strip() and len(str(data.get("pitch_deck_url")).strip()) > 5)
    summary = str(data.get("business_summary") or "").lower()
    summary_words = len(summary.split()) if summary else 0

    score_signals = 0

    if revenue >= 5000000:
        score_signals += 6
    elif revenue >= 2000000:
        score_signals += 5
    elif revenue >= 500000:
        score_signals += 4
    elif revenue > 0:
        score_signals += 3

    if team_size >= 10:
        score_signals += 4
    elif team_size >= 5:
        score_signals += 3
    elif team_size >= 2:
        score_signals += 2

    if has_website:
        score_signals += 2
    if has_dpiit:
        score_signals += 2
    if has_pitch:
        score_signals += 1

    if summary_words > 40:
        score_signals += 2
    elif summary_words > 15:
        score_signals += 1

    if score_signals <= 2:
        return STAGE_IDEATION
    elif score_signals <= 5:
        return STAGE_PROTOTYPE
    elif score_signals <= 8:
        return STAGE_MVP
    elif score_signals <= 11:
        return STAGE_REVENUE
    elif score_signals <= 14:
        return STAGE_GROWTH
    else:
        return STAGE_SEED


# ─── Stage-Specific Scoring Engines (each max 40 raw → normalized 0-100) ──

def _score_ideation(startup: Dict[str, Any]) -> Tuple[float, Dict[str, Any]]:
    scores = {}
    summary = str(startup.get("business_summary") or "").lower()
    sector = str(startup.get("sector") or "").lower()
    competitors = str(startup.get("competitors") or "").lower()

    word_count = len(summary.split())
    if word_count > 40:
        scores["idea_quality"] = 10
    elif word_count > 20:
        scores["idea_quality"] = 7
    elif word_count > 5:
        scores["idea_quality"] = 4
    else:
        scores["idea_quality"] = 1

    problem_keywords = ["problem", "solve", "challenge", "issue", "gap", "need", "pain"]
    hits = sum(1 for kw in problem_keywords if kw in summary)
    scores["problem_clarity"] = min(8, 2 + hits * 2)

    high_market = ["health", "agri", "climate", "ai", "fintech", "edtech", "saas"]
    if any(m in sector for m in high_market):
        scores["market_potential"] = 8
    elif sector and len(sector) > 2:
        scores["market_potential"] = 5
    else:
        scores["market_potential"] = 2

    size = int(startup.get("team_size") or 1)
    if size >= 3:
        scores["team"] = 6
    elif size == 2:
        scores["team"] = 4
    else:
        scores["team"] = 2

    if competitors and len(competitors) > 5 and "none" not in competitors:
        scores["competitor_awareness"] = 4
    else:
        scores["competitor_awareness"] = 1

    extras = 0
    if str(startup.get("website") or "").strip():
        extras += 1
    if startup.get("dpiit"):
        extras += 2
    if str(startup.get("pitch_deck_url") or "").strip():
        extras += 1
    scores["extras"] = extras

    total = sum(scores.values())
    normalized = (total / 40.0) * 100.0
    return round(normalized, 1), scores


def _score_prototype(startup: Dict[str, Any]) -> Tuple[float, Dict[str, Any]]:
    scores = {}
    summary = str(startup.get("business_summary") or "").lower()

    website = str(startup.get("website") or "").strip()
    if website and website.startswith("http"):
        scores["product_presence"] = 8
    elif website:
        scores["product_presence"] = 5
    else:
        scores["product_presence"] = 1

    size = int(startup.get("team_size") or 1)
    if size >= 3:
        scores["team"] = 6
    elif size == 2:
        scores["team"] = 4
    else:
        scores["team"] = 2

    tech_keywords = ["built", "developed", "prototype", "working", "demo", "architecture", "technology", "platform"]
    hits = sum(1 for kw in tech_keywords if kw in summary)
    word_count = len(summary.split())
    if hits >= 2 and word_count > 20:
        scores["technical_detail"] = 8
    elif hits >= 1 or word_count > 15:
        scores["technical_detail"] = 5
    else:
        scores["technical_detail"] = 2

    scores["dpiit"] = 4 if startup.get("dpiit") else 0

    sector = str(startup.get("sector") or "").lower()
    deep_tech = ["ai", "ml", "iot", "blockchain", "robotics", "biotech", "deeptech", "hardware"]
    if any(t in sector for t in deep_tech):
        scores["sector"] = 6
    elif sector and len(sector) > 2:
        scores["sector"] = 4
    else:
        scores["sector"] = 2

    pitch = str(startup.get("pitch_deck_url") or "").strip()
    scores["pitch_deck"] = 4 if pitch and len(pitch) > 5 else 0

    rev = float(startup.get("revenue") or 0.0)
    scores["early_revenue"] = 4 if rev > 0 else 2

    total = sum(scores.values())
    normalized = (total / 40.0) * 100.0
    return round(normalized, 1), scores


def _score_mvp(startup: Dict[str, Any]) -> Tuple[float, Dict[str, Any]]:
    scores = {}
    summary = str(startup.get("business_summary") or "").lower()

    website = str(startup.get("website") or "").strip()
    if website and website.startswith("http"):
        scores["website"] = 8
    elif website:
        scores["website"] = 5
    else:
        scores["website"] = 1

    size = int(startup.get("team_size") or 1)
    if size >= 4:
        scores["team"] = 6
    elif size >= 2:
        scores["team"] = 4
    else:
        scores["team"] = 2

    scores["dpiit"] = 5 if startup.get("dpiit") else 0

    word_count = len(summary.split())
    if word_count > 30:
        scores["summary_depth"] = 6
    elif word_count > 15:
        scores["summary_depth"] = 4
    else:
        scores["summary_depth"] = 1

    pitch = str(startup.get("pitch_deck_url") or "").strip()
    scores["pitch_deck"] = 5 if pitch and len(pitch) > 5 else 0

    sector = str(startup.get("sector") or "").lower()
    scores["sector"] = 5 if sector and len(sector) > 2 else 2

    rev = float(startup.get("revenue") or 0.0)
    scores["revenue"] = 5 if rev > 0 else 3

    total = sum(scores.values())
    normalized = (total / 40.0) * 100.0
    return round(normalized, 1), scores


def _score_revenue(startup: Dict[str, Any]) -> Tuple[float, Dict[str, Any]]:
    scores = {}

    rev = float(startup.get("revenue") or 0.0)
    if rev >= 2000000:
        scores["revenue"] = 12
    elif rev >= 500000:
        scores["revenue"] = 9
    elif rev > 0:
        scores["revenue"] = 6
    else:
        scores["revenue"] = 0

    summary = str(startup.get("business_summary") or "").lower()
    growth_kw = ["growing", "growth", "increasing", "month on month", "mrr", "arr", "recurring", "retention", "repeat"]
    hits = sum(1 for kw in growth_kw if kw in summary)
    scores["growth_signals"] = min(5, hits * 2)

    size = int(startup.get("team_size") or 1)
    if size >= 5:
        scores["team"] = 5
    elif size >= 3:
        scores["team"] = 4
    elif size >= 2:
        scores["team"] = 3
    else:
        scores["team"] = 1

    website = str(startup.get("website") or "").strip()
    scores["website"] = 5 if website and website.startswith("http") else 0

    scores["dpiit"] = 4 if startup.get("dpiit") else 0

    pitch = str(startup.get("pitch_deck_url") or "").strip()
    scores["pitch_deck"] = 4 if pitch and len(pitch) > 5 else 0

    sector = str(startup.get("sector") or "").lower()
    scores["sector"] = 5 if sector and len(sector) > 2 else 2

    total = sum(scores.values())
    normalized = (total / 40.0) * 100.0
    return round(normalized, 1), scores


def _score_growth(startup: Dict[str, Any]) -> Tuple[float, Dict[str, Any]]:
    scores = {}

    rev = float(startup.get("revenue") or 0.0)
    if rev >= 10000000:
        scores["revenue_magnitude"] = 10
    elif rev >= 5000000:
        scores["revenue_magnitude"] = 8
    elif rev >= 2000000:
        scores["revenue_magnitude"] = 6
    elif rev > 0:
        scores["revenue_magnitude"] = 3
    else:
        scores["revenue_magnitude"] = 0

    size = int(startup.get("team_size") or 1)
    if size >= 15:
        scores["team_scale"] = 8
    elif size >= 10:
        scores["team_scale"] = 7
    elif size >= 5:
        scores["team_scale"] = 5
    elif size >= 3:
        scores["team_scale"] = 3
    else:
        scores["team_scale"] = 1

    website = str(startup.get("website") or "").strip()
    scores["website"] = 5 if website and website.startswith("http") else 0

    scores["dpiit"] = 4 if startup.get("dpiit") else 0

    summary = str(startup.get("business_summary") or "").lower()
    scale_kw = ["scale", "scaling", "expand", "expansion", "market", "revenue", "growth", "traction", "users", "customers", "retention"]
    hits = sum(1 for kw in scale_kw if kw in summary)
    scores["scaling_signals"] = min(5, 1 + hits)

    sector = str(startup.get("sector") or "").lower()
    scores["sector"] = 4 if sector and len(sector) > 2 else 1

    pitch = str(startup.get("pitch_deck_url") or "").strip()
    scores["pitch_deck"] = 4 if pitch and len(pitch) > 5 else 0

    total = sum(scores.values())
    normalized = (total / 40.0) * 100.0
    return round(normalized, 1), scores


def _score_seed(startup: Dict[str, Any]) -> Tuple[float, Dict[str, Any]]:
    scores = {}

    rev = float(startup.get("revenue") or 0.0)
    if rev >= 10000000:
        scores["revenue"] = 8
    elif rev >= 5000000:
        scores["revenue"] = 7
    elif rev >= 2000000:
        scores["revenue"] = 5
    elif rev > 0:
        scores["revenue"] = 3
    else:
        scores["revenue"] = 0

    pitch = str(startup.get("pitch_deck_url") or "").strip()
    if pitch and len(pitch) > 10:
        scores["pitch_deck"] = 8
    elif pitch and len(pitch) > 5:
        scores["pitch_deck"] = 5
    else:
        scores["pitch_deck"] = 0

    website = str(startup.get("website") or "").strip()
    scores["website"] = 5 if website and website.startswith("http") else 0

    size = int(startup.get("team_size") or 1)
    if size >= 10:
        scores["team"] = 5
    elif size >= 5:
        scores["team"] = 4
    elif size >= 3:
        scores["team"] = 3
    else:
        scores["team"] = 1

    scores["dpiit"] = 4 if startup.get("dpiit") else 0

    summary = str(startup.get("business_summary") or "").lower()
    invest_kw = ["investment", "raise", "funding", "seed", "series", "valuation", "investor", "capital", "financial", "projection"]
    hits = sum(1 for kw in invest_kw if kw in summary)
    scores["investment_readiness"] = min(5, hits)

    sector = str(startup.get("sector") or "").lower()
    high_value = ["ai", "saas", "fintech", "healthtech", "deeptech", "climate", "biotech"]
    if any(s in sector for s in high_value):
        scores["sector"] = 5
    elif sector and len(sector) > 2:
        scores["sector"] = 3
    else:
        scores["sector"] = 1

    total = sum(scores.values())
    normalized = (total / 40.0) * 100.0
    return round(normalized, 1), scores


_STAGE_SCORERS = {
    STAGE_IDEATION:  _score_ideation,
    STAGE_PROTOTYPE: _score_prototype,
    STAGE_MVP:       _score_mvp,
    STAGE_REVENUE:   _score_revenue,
    STAGE_GROWTH:    _score_growth,
    STAGE_SEED:      _score_seed,
}

_STAGE_PRIORITY_THRESHOLDS = {
    STAGE_IDEATION:  (60, 35),
    STAGE_PROTOTYPE: (55, 30),
    STAGE_MVP:       (55, 30),
    STAGE_REVENUE:   (50, 25),
    STAGE_GROWTH:    (45, 20),
    STAGE_SEED:      (45, 20),
}

_COHORT_PRIORITY_THRESHOLDS = {
    STAGE_IDEATION:  (70, 40),
    STAGE_PROTOTYPE: (65, 35),
    STAGE_MVP:       (65, 35),
    STAGE_REVENUE:   (60, 30),
    STAGE_GROWTH:    (55, 25),
    STAGE_SEED:      (55, 25),
}


load_dotenv()

# Setup Fernet Encryption Key
ENCRYPTION_KEY = os.environ.get("ENCRYPTION_KEY")
if not ENCRYPTION_KEY:
    # Generate and save to env
    key = Fernet.generate_key().decode()
    env_path = os.path.join(os.path.dirname(os.path.dirname(os.path.abspath(__file__))), ".env")
    if os.path.exists(env_path):
        try:
            with open(env_path, "a") as f:
                f.write(f"\nENCRYPTION_KEY={key}\n")
        except Exception as e:
            print(f"Error writing ENCRYPTION_KEY to .env: {e}")
    os.environ["ENCRYPTION_KEY"] = key
    ENCRYPTION_KEY = key

cipher = Fernet(ENCRYPTION_KEY.encode())

def encrypt_val(val: Any) -> str:
    if val is None:
        return ""
    val_str = str(val).strip()
    if not val_str:
        return ""
    return cipher.encrypt(val_str.encode()).decode()

def decrypt_val(val: str) -> str:
    if not val:
        return ""
    try:
        return cipher.decrypt(val.encode()).decode()
    except Exception:
        return val # Fallback if decryption fails or if it wasn't encrypted

def clean_revenue(rev_str: Any) -> float:
    if rev_str is None:
        return 0.0
    s = str(rev_str).strip().lower()
    if not s or s in ["none", "na", "n/a", "no", "nil", "0", "zero"]:
        return 0.0
    
    # Remove currency symbols, spaces, commas
    s = re.sub(r"[₹$,\s]", "", s)
    
    multiplier = 1.0
    if "lakh" in s or "l" in s:
        multiplier = 100000.0
        s = re.sub(r"[a-z]", "", s)
    elif "crore" in s or "cr" in s:
        multiplier = 10000000.0
        s = re.sub(r"[a-z]", "", s)
        
    try:
        # Extract first decimal number found
        match = re.search(r"[-+]?\d*\.\d+|\d+", s)
        if match:
            return float(match.group()) * multiplier
        return 0.0
    except ValueError:
        return 0.0

def clean_team_size(team_str: Any) -> int:
    if team_str is None:
        return 1
    s = str(team_str).strip()
    if not s:
        return 1
    # Extract first integer
    match = re.search(r"\d+", s)
    if match:
        return int(match.group())
    return 1

def clean_dpiit(dpiit_str: Any) -> Tuple[bool, str]:
    if dpiit_str is None:
        return False, ""
    s = str(dpiit_str).strip()
    s_lower = s.lower()
    if not s_lower or s_lower in ["no", "none", "na", "n/a", "false"]:
        return False, ""
    
    match = re.search(r"\d+", s)
    dpiit_num = match.group() if match else s
    return True, dpiit_num

def map_excel_headers(headers: List[str]) -> Dict[str, int]:
    mapping = {}
    header_patterns = {
        "email": [r"email", r"e-mail", r"mail"],
        "name": [r"founder.*name", r"applicant.*name", r"name\s*\(", r"^name$", r"contact\s*person"],
        "mobile": [r"mobile\s*number", r"^mobile$", r"phone", r"contact\s*no"],
        "alternet_mobile": [r"alternet", r"alternative.*mobile", r"alt.*phone"],
        "dob": [r"date\s*of\s*birth", r"dob"],
        "gender": [r"gender"],
        "address": [r"address", r"location"],
        "city_state": [r"city\s*&\s*state", r"city", r"state", r"region"],
        "highest_qualification": [r"highest\s*qualification", r"qualification", r"degree"],
        "school_university": [r"school", r"college", r"university", r"institution"],
        "how_found_out": [r"how\s*did\s*you\s*find\s*out", r"source"],
        "startup_name": [r"name\s*of\s*startup", r"company/idea\s*name", r"startup.*name", r"incubator\s*name", r"hub\s*name", r"center\s*name", r"org.*name", r"tbi\s*name", r"entity\s*name", r"company\s*name", r"title"],
        "year_established": [r"year\s*of\s*establishment", r"established", r"founding\s*year"],
        "sector": [r"sector", r"domain", r"industry", r"focus\s*area", r"vertical"],
        "team_size": [r"team\s*members", r"team\s*size", r"capacity", r"number\s*of\s*startups"],
        "company_registered_name": [r"company\s*name\s*\(", r"registered\s*name"],
        "comments": [r"comments\s*or\s*questions", r"remarks", r"notes"],
        "dpiit_registered": [r"dpiit\s*registered", r"dpiit", r"recognized", r"registration\s*status"],
        "website": [r"website", r"share\s*the\s*url", r"website.*url", r"portal\s*url", r"url"],
        "applied_other": [r"applied\s*to\s*any\s*other\s*incubator", r"applied\s*elsewhere"],
        "litigation": [r"litigation", r"legal\s*disputes"],
        "business_summary": [r"summary\s*of\s*your\s*business\s*idea", r"business\s*summary", r"description", r"overview", r"about", r"details"],
        "competitors": [r"competitors", r"competition"],
        "revenue": [r"revenue", r"turnover", r"grants\s*received", r"funding\s*raised", r"budget"],
        "pitch_deck_url": [r"presentation", r"pitch\s*deck", r"deck\s*url", r"attachment"],
        "stage": [r"stage\s*of\s*startup", r"^stage$", r"maturity", r"level", r"type"],
        "legal_entity": [r"legal\s*entity\s*type", r"entity\s*type", r"constitution"],
        "logo_url": [r"company\s*logo", r"logo"],
        "applying_for": [r"what\s*are\s*you\s*applying\s*for", r"program"]
    }
    
    for key, patterns in header_patterns.items():
        for i, h in enumerate(headers):
            if h is None:
                continue
            h_str = str(h).strip().lower()
            for pat in patterns:
                if re.search(pat, h_str):
                    mapping[key] = i
                    break
            if key in mapping:
                break
                
    return mapping

def extract_dynamic_rows_and_headers(sheet) -> Tuple[List[str], List[Dict[str, Any]], Dict[str, int]]:
    headers = []
    first_row = sheet[1]
    for cell in first_row:
        val = str(cell.value).strip() if cell.value is not None else ""
        headers.append(val)

    header_map = map_excel_headers(headers)
    
    # Inspect sample row values to auto-classify email and phone columns
    sample_rows = []
    for r_idx in range(2, min(sheet.max_row + 1, 6)):
        r_vals = [sheet.cell(row=r_idx, column=c_idx).value for c_idx in range(1, len(headers) + 1)]
        if any(r_vals):
            sample_rows.append(r_vals)
            
    email_col_indices = set()
    phone_col_indices = set()
    
    for c_idx in range(len(headers)):
        for r_vals in sample_rows:
            if c_idx < len(r_vals) and r_vals[c_idx] is not None:
                val_str = str(r_vals[c_idx]).strip()
                if "@" in val_str and "." in val_str and not val_str.startswith("http"):
                    email_col_indices.add(c_idx)
                elif re.search(r"^[+]?\d{10,12}$", val_str.replace(" ", "").replace("-", "")):
                    phone_col_indices.add(c_idx)

    # Automatically set email / mobile in header_map if detected
    if "email" not in header_map and email_col_indices:
        header_map["email"] = list(email_col_indices)[0]

    if "mobile" not in header_map and phone_col_indices:
        header_map["mobile"] = list(phone_col_indices)[0]

    # Identify primary Name column: MUST NOT be an email or phone column
    name_col_idx = None
    for key in ["startup_name", "company_registered_name", "name"]:
        if key in header_map and header_map[key] not in email_col_indices and header_map[key] not in phone_col_indices:
            name_col_idx = header_map[key]
            break
            
    if name_col_idx is None:
        # Fallback 1: Header with name keywords excluding email/phone
        for idx, h in enumerate(headers):
            if idx in email_col_indices or idx in phone_col_indices:
                continue
            h_lower = h.lower()
            if any(term in h_lower for term in ["name", "title", "startup", "company", "incubator", "entity", "hub", "center", "organization"]):
                name_col_idx = idx
                break
                
    if name_col_idx is None:
        # Fallback 2: First non-email, non-phone column with non-empty header or data
        for idx, h in enumerate(headers):
            if idx not in email_col_indices and idx not in phone_col_indices:
                name_col_idx = idx
                break
                
    if name_col_idx is None and len(headers) > 0:
        name_col_idx = 0

    rows_data = []
    for row_idx in range(2, sheet.max_row + 1):
        row_values = [sheet.cell(row=row_idx, column=col_idx).value for col_idx in range(1, len(headers) + 1)]
        if not any(row_values):
            continue
            
        raw_dict = {}
        for h_idx, h_name in enumerate(headers):
            key_name = h_name if h_name else f"Column_{h_idx + 1}"
            val = row_values[h_idx] if h_idx < len(row_values) else None
            raw_dict[key_name] = "" if val is None else str(val).strip()

        raw_entity_val = str(row_values[name_col_idx]).strip() if (name_col_idx is not None and name_col_idx < len(row_values) and row_values[name_col_idx] is not None) else f"Entry #{row_idx - 1}"
        
        # If entity name column happens to be an email address, parse out clean human name
        if "@" in raw_entity_val and "." in raw_entity_val:
            prefix = raw_entity_val.split("@")[0]
            cleaned_name = re.sub(r"\d+", "", prefix).replace(".", " ").replace("_", " ").title().strip()
            entity_name = cleaned_name if cleaned_name else prefix.title()
        else:
            entity_name = raw_entity_val

        if not entity_name or entity_name.lower() in ["none", "null", "n/a"]:
            entity_name = f"Entry #{row_idx - 1}"

        rows_data.append({
            "row_idx": row_idx,
            "entity_name": entity_name,
            "raw_data": raw_dict,
            "row_values": row_values
        })

    return headers, rows_data, header_map

def evaluate_dynamic_features(raw_data: Dict[str, str], headers: List[str]) -> Tuple[float, Dict[str, float], List[str], List[str]]:
    """
    Evaluates ANY random excel columns and returns dynamic score breakdown, strengths, and weaknesses.
    """
    feature_scores = {}
    strengths = []
    weaknesses = []
    
    total_dynamic_score = 0.0
    evaluated_count = 0
    
    for h_name, val_str in raw_data.items():
        if not h_name or not val_str:
            continue
            
        h_lower = h_name.lower()
        val_lower = val_str.lower()
        col_score = 0.0
        
        # 1. Check if column value is numeric
        num_match = re.search(r"[-+]?\d*\.\d+|\d+", val_str.replace(",", ""))
        if num_match:
            try:
                num_val = float(num_match.group())
                if any(kw in h_lower for kw in ["score", "rating", "rank", "marks", "points", "grade"]):
                    if num_val <= 10:
                        col_score = min(10.0, num_val)
                    elif num_val <= 100:
                        col_score = min(10.0, (num_val / 100.0) * 10.0)
                    else:
                        col_score = 8.0
                elif any(kw in h_lower for kw in ["revenue", "funding", "grant", "budget", "capital", "turnover"]):
                    cleaned_rev = clean_revenue(val_str)
                    if cleaned_rev > 2000000:
                        col_score = 10.0
                    elif cleaned_rev > 500000:
                        col_score = 7.5
                    elif cleaned_rev > 0:
                        col_score = 5.0
                    else:
                        col_score = 2.0
                elif any(kw in h_lower for kw in ["team", "member", "staff", "employee", "startups", "capacity"]):
                    team_n = clean_team_size(val_str)
                    col_score = 10.0 if team_n >= 5 else (7.0 if team_n >= 2 else 4.0)
                else:
                    col_score = 6.0 if num_val > 0 else 3.0
                    
                feature_scores[h_name] = round(col_score * 10.0, 1)
                total_dynamic_score += col_score
                evaluated_count += 1
                if col_score >= 7.5:
                    strengths.append(f"Strong {h_name}: {val_str}")
                continue
            except Exception:
                pass

        # 2. Check affirmative / boolean signals
        if val_lower in ["yes", "true", "dpiit", "active", "certified", "registered", "granted", "approved", "high", "top"]:
            col_score = 9.0
            feature_scores[h_name] = 90.0
            total_dynamic_score += col_score
            evaluated_count += 1
            strengths.append(f"Verified {h_name} ({val_str})")
        elif val_lower in ["no", "false", "pending", "none", "n/a", "nil"]:
            col_score = 2.0
            feature_scores[h_name] = 20.0
            total_dynamic_score += col_score
            evaluated_count += 1
            weaknesses.append(f"Missing/Unconfirmed {h_name}")
        elif len(val_str) > 10: # Text fields
            word_count = len(val_str.split())
            if word_count > 25:
                col_score = 9.0
                strengths.append(f"Detailed {h_name}")
            elif word_count > 8:
                col_score = 6.5
            else:
                col_score = 4.5
            feature_scores[h_name] = round(col_score * 10.0, 1)
            total_dynamic_score += col_score
            evaluated_count += 1

    if evaluated_count > 0:
        avg_score = (total_dynamic_score / (evaluated_count * 10.0)) * 100.0
    else:
        avg_score = 50.0

    return round(avg_score, 1), feature_scores, strengths[:3], weaknesses[:3]

def evaluate_rules(startup: Dict[str, Any], stage_category: Optional[str] = None) -> Tuple[float, Dict[str, Any]]:
    if not stage_category:
        stage_category = classify_startup_stage(startup)

    scorer = _STAGE_SCORERS.get(stage_category)
    if scorer:
        normalized, scores = scorer(startup)
    else:
        # Fallback to legacy generic scoring
        scores = {}
        rev = startup.get("revenue", 0.0)
        scores["revenue"] = 10 if rev > 2000000 else (8 if rev > 500000 else (5 if rev > 0 else 0))
        scores["dpiit"] = 5 if startup.get("dpiit", False) else 0
        size = startup.get("team_size", 1)
        scores["team_size"] = 5 if size >= 5 else (4 if size >= 2 else 2)
        website = str(startup.get("website", "")).strip()
        scores["website"] = 5 if website and (website.startswith("http") or "." in website) else 0
        pitch = str(startup.get("pitch_deck_url", "")).strip()
        scores["pitch_deck"] = 5 if pitch and ("http" in pitch or len(pitch) > 5) else 0
        scores["stage"] = 5
        total_score = sum(scores.values())
        normalized = (total_score / 40.0) * 100.0
        normalized, scores = round(normalized, 1), scores

    scores["stage_category"] = stage_category
    return normalized, scores

def evaluate_advanced_heuristics(startup: Dict[str, Any], stage_category: Optional[str] = None) -> Dict[str, Any]:
    if not stage_category:
        stage_category = classify_startup_stage(startup)

    summary = str(startup.get("business_summary") or "").lower()
    sector = str(startup.get("sector") or "").lower()
    competitors = str(startup.get("competitors") or "").lower()
    team_size = int(startup.get("team_size") or 1)
    revenue = float(startup.get("revenue") or 0.0)
    stage = str(startup.get("stage") or "").lower()
    website = str(startup.get("website") or "").strip()
    dpiit = startup.get("dpiit", False)
    
    innovation = 10
    market = 10
    scalability = 10
    execution = 10
    problem = 10
    
    strengths = []
    weaknesses = []
    
    # 1. Market & Scalability
    if "tech" in sector or "ai" in sector or "ml" in sector or "software" in sector or "saas" in sector:
        market += 4
        scalability += 5
        strengths.append("High-growth potential tech sector")
    elif "health" in sector or "medical" in sector or "agri" in sector or "climate" in sector:
        market += 5
        scalability += 3
        strengths.append("Addresses critical, large-scale sector")
    else:
        market += 2
        scalability += 2
        
    if competitors and len(competitors) > 5 and "none" not in competitors:
        market += 3
        strengths.append("Aware of competitive landscape")
    else:
        weaknesses.append("Lack of competitor awareness or no clear competitors mentioned")
        
    # 2. Innovation & Problem
    if "ai " in summary or "ml " in summary or "patent" in summary or "proprietary" in summary:
        innovation += 6
        strengths.append("Strong focus on innovation/IP")
    elif len(summary.split()) > 30:
        innovation += 3
        problem += 4
        strengths.append("Detailed business summary and problem statement")
    else:
        weaknesses.append("Vague or overly brief business summary")
        
    if dpiit:
        innovation += 2
        problem += 3
        strengths.append("DPIIT Recognized (Validated startup)")
        
    # 3. Execution (Traction & Readiness)
    if revenue > 0:
        execution += 8
        strengths.append("Demonstrated revenue traction")
    elif "revenue" in stage or "traction" in stage or "growth" in stage:
        execution += 6
        strengths.append("Reported traction stage")
    elif "mvp" in stage or "prototype" in stage:
        execution += 4
        strengths.append("Product development is underway")
    else:
        weaknesses.append("Early stage with unproven execution")
        
    if team_size > 1:
        execution += 2
        strengths.append("Has a co-founding team/multiple members")
    else:
        weaknesses.append("Solo founder risk")
        
    if website:
        scalability += 2
    else:
        weaknesses.append("Lacks digital presence (no website)")
        
    # Cap scores at 20
    innovation = min(20, innovation)
    market = min(20, market)
    scalability = min(20, scalability)
    execution = min(20, execution)
    problem = min(20, problem)
    
    total = innovation + market + scalability + execution + problem
    
    if total >= 80:
        recommendation = "Highly Recommended"
    elif total >= 65:
        recommendation = "Recommended"
    elif total >= 50:
        recommendation = "Consider with Reservations"
    else:
        recommendation = "Not Recommended"
        
    if not strengths:
        strengths.append("No distinct strengths identified from provided data")
    if not weaknesses:
        weaknesses.append("No obvious weaknesses detected")
        
    return {
        "innovation": innovation,
        "market": market,
        "scalability": scalability,
        "execution": execution,
        "problem": problem,
        "strengths": strengths[:3],
        "weaknesses": weaknesses[:3],
        "recommendation": recommendation,
        "llm_score": float(total),
        "stage_category": stage_category,
    }

def compute_similarity_matrix(startups: List[Dict[str, Any]]) -> List[Dict[str, Any]]:
    summaries = [s.get("business_summary", "") for s in startups]
    n = len(startups)
    
    for s in startups:
        s["similarity_matches"] = []
        
    if n <= 1:
        return startups
        
    try:
        from sklearn.feature_extraction.text import TfidfVectorizer
        from sklearn.metrics.pairwise import cosine_similarity
        
        valid_indices = [i for i, text in enumerate(summaries) if text.strip()]
        if len(valid_indices) <= 1:
            return startups
            
        valid_summaries = [summaries[i] for i in valid_indices]
        
        vectorizer = TfidfVectorizer(stop_words='english')
        tfidf_matrix = vectorizer.fit_transform(valid_summaries)
        sim_matrix = cosine_similarity(tfidf_matrix, tfidf_matrix)
        
        for i in range(len(valid_indices)):
            orig_idx = valid_indices[i]
            for j in range(len(valid_indices)):
                if i == j:
                    continue
                score = sim_matrix[i][j]
                if score >= 0.4: # Log significant matches
                    other_idx = valid_indices[j]
                    startups[orig_idx]["similarity_matches"].append({
                        "startup_name": startups[other_idx].get("startup_name", "Unknown"),
                        "similarity_score": round(float(score) * 100, 1)
                    })
    except Exception as e:
        print(f"Error computing TF-IDF similarity, using fallback string overlap: {e}")
        for i in range(n):
            words_i = set(re.findall(r"\w+", summaries[i].lower()))
            if not words_i:
                continue
            for j in range(n):
                if i == j:
                    continue
                words_j = set(re.findall(r"\w+", summaries[j].lower()))
                if not words_j:
                    continue
                intersection = words_i.intersection(words_j)
                union = words_i.union(words_j)
                score = len(intersection) / len(union) if union else 0.0
                if score >= 0.3:
                    startups[i]["similarity_matches"].append({
                        "startup_name": startups[j].get("startup_name", "Unknown"),
                        "similarity_score": round(score * 100, 1)
                    })
                    
    for s in startups:
        s["similarity_matches"].sort(key=lambda x: x["similarity_score"], reverse=True)
        
    return startups
